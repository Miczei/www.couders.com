from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

F = "Arial"
NAG = PatternFill("solid", fgColor="0F6B5F")
NAG_F = Font(name=F, bold=True, color="FFFFFF", size=10)
TXT = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
TYT = Font(name=F, size=14, bold=True, color="0B4D44")
PRZYK = Font(name=F, size=10, italic=True, color="8A6516")
ZOLTY = PatternFill("solid", fgColor="FFF3CD")
SZARY = PatternFill("solid", fgColor="F2F6F4")
CIEN = Border(bottom=Side(style="thin", color="D0D9D4"))
GORA = Alignment(vertical="top", wrap_text=True)

wb = Workbook()

def naglowki(ws, kols, szer):
    ws.append(kols)
    for i, k in enumerate(kols, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = NAG, NAG_F
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = szer[i-1]
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

def wiersze(ws, dane, przyklad_pierwszy=False):
    for n, r in enumerate(dane):
        ws.append(r)
        rw = ws.max_row
        for i in range(1, len(r) + 1):
            c = ws.cell(row=rw, column=i)
            c.font = PRZYK if (przyklad_pierwszy and n == 0) else TXT
            c.alignment, c.border = GORA, CIEN
        if przyklad_pierwszy and n == 0:
            for i in range(1, len(r) + 1):
                ws.cell(row=rw, column=i).fill = ZOLTY

# ============ INSTRUKCJA ============
ws = wb.active
ws.title = "instrukcja"
ws.column_dimensions["A"].width = 118
ws.sheet_view.showGridLines = False

tekst = [
 ("Lux Dom, baza agenta głosowego. Couders.", "tyt"),
 ("", ""),
 ("Wersja demonstracyjna. Dane dwóch nieruchomości są PRZYKŁADOWE i służą do pokazu.", ""),
 ("Przed wdrożeniem podmienia się je na realny portfel Lux Dom.", ""),
 ("", ""),
 ("ZAKŁADKI", "b"),
 ("• nieruchomosci — obsługiwane wspólnoty. Kolumna klucz to wartość, po której n8n rozpoznaje", ""),
 ("   adres z rozmowy. Kolumna adres_glosowy to gotowa forma do wymówienia, już odmieniona", ""),
 ("   (\"Sarmackiej dwanaście\"), bo agent czyta ją na głos bez przetwarzania.", ""),
 ("   Kolumna zgloszen_w_bazie liczy się formułą, więc zostaje aktualna po każdym zgłoszeniu.", ""),
 ("• lokale — numery rozpoznawane po caller ID. Agent wita takiego rozmówcę po imieniu", ""),
 ("   i od razu zna jego nieruchomość oraz lokal. TU WPISZ NUMERY OSÓB, KTÓRE BĘDĄ NA SPOTKANIU.", ""),
 ("• zgloszenia — tu agent zapisuje sprawy. Wiersz 2 to PRZYKŁAD formatu, skasuj go przed", ""),
 ("   pierwszym prawdziwym testem. NAGŁÓWKÓW NIE RUSZAJ ani nie zmieniaj kolejności:", ""),
 ("   mapowanie w n8n jest defineBelow i szuka kolumn po nazwie.", ""),
 ("• baza_wiedzy — treści, z których agent odpowiada. Kolumna odpowiedz_glosowa to zdanie", ""),
 ("   gotowe do przeczytania w telefonie, bez skrótów i bez cyfr zapisanych znakami.", ""),
 ("• slowniki — dozwolone wartości pól w narzędziach. Agent musi trzymać się tej listy,", ""),
 ("   inaczej n8n dostanie kategorię, której nie zna.", ""),
 ("", ""),
 ("JAK PODPIĄĆ", "b"),
 ("1. Wgraj ten plik na Dysk Google i otwórz go w Arkuszach.", ""),
 ("2. Skopiuj ID arkusza z paska adresu.", ""),
 ("3. Wklej je w węźle Google Sheets w workflow n8n, w miejsce PODMIEN_ID_ARKUSZA.", ""),
 ("4. W tym samym węźle zostaw operację Append or Update i kolumnę dopasowania id_rozmowy.", ""),
 ("", ""),
 ("DLACZEGO id_rozmowy JEST WAŻNE", "b"),
 ("Zgłoszenie powstaje w trakcie rozmowy, a transkrypt i streszczenie dochodzą dopiero po jej", ""),
 ("zakończeniu, z post-call webhooka. Oba trafiają do TEGO SAMEGO wiersza, dopasowane po", ""),
 ("id_rozmowy. Jeśli ta kolumna będzie pusta, powstaną dwa osobne wiersze zamiast jednego.", ""),
 ("", ""),
 ("CZEGO AGENT NIE ZAPISUJE", "b"),
 ("• Duplikatów. Czterdzieści telefonów o tej samej awarii daje jeden wiersz i czterdzieści", ""),
 ("   powiadomień po naprawie, a nie czterdzieści zleceń dla technika.", ""),
 ("• Sald i kwot zadłużenia. Tych danych agent nie podaje przez telefon i nie zapisuje tutaj.", ""),
 ("", ""),
 ("DO POTWIERDZENIA U KLIENTA", "b"),
 ("• Godziny pracy biura i zakres dyżuru poza godzinami. Wpisane wartości są założeniem.", ""),
 ("• Numery kont wspólnot. W bazie trzymamy tylko końcówkę, bo agent nie czyta pełnego numeru.", ""),
 ("• Ścieżki alarmowe: kto i pod jakim numerem odbiera awarię krytyczną dla każdej nieruchomości.", ""),
]
for t, styl in tekst:
    ws.append([t])
    c = ws.cell(row=ws.max_row, column=1)
    c.font = TYT if styl == "tyt" else (BOLD if styl == "b" else TXT)
    c.alignment = Alignment(vertical="center")

# ============ NIERUCHOMOSCI ============
ws = wb.create_sheet("nieruchomosci")
naglowki(ws, ["klucz","nazwa_wspolnoty","adres","adres_glosowy","miasto","lokali","administrator",
              "administrator_telefon","administrator_mail","konto_koncowka","odczyty","odpady",
              "zgloszen_w_bazie","uwaga"],
         [16,26,18,22,12,8,18,18,26,14,22,26,15,34])
wiersze(ws, [
 ["sarmacka 12","Wspólnota Sarmacka 12","Sarmacka 12","Sarmackiej dwanaście","Warszawa Wilanów",68,
  "Katarzyna Nowak","22 292 19 15","k.nowak@lux-dom.pl","0001","do 10 dnia miesiąca",
  "zmieszane pon i czw, segregacja wt","","hala garażowa, winda, domofon, węzeł cieplny"],
 ["klimczaka 5","Wspólnota Klimczaka 5","Klimczaka 5","Klimczaka pięć","Warszawa Wilanów",124,
  "Piotr Zieliński","22 292 19 15","p.zielinski@lux-dom.pl","0002","do 10 dnia miesiąca",
  "zmieszane wt i pt, segregacja śr","","dwie klatki, dwie windy, brama na pilota, wentylacja mechaniczna"],
])
for r in (2, 3):
    ws.cell(row=r, column=13).value = f'=COUNTIF(zgloszenia!C:C,C{r})'
    ws.cell(row=r, column=13).font = TXT

# ============ LOKALE ============
ws = wb.create_sheet("lokale")
naglowki(ws, ["telefon","imie_nazwisko","zwrot_glosowy","forma","klucz_nieruchomosci","lokal","rola","uwaga"],
         [16,24,18,8,18,10,16,42])
wiersze(ws, [
 ["+48600100200","Marek Wiśniewski (PRZYKŁAD, podmień)","panie Marku","Pan","sarmacka 12","34","wlasciciel",
  "PODMIEŃ na numer osoby, która będzie na spotkaniu. Format +48XXXXXXXXX"],
 ["+48600100201","Anna Kowalska (PRZYKŁAD, podmień)","pani Anno","Pani","klimczaka 5","12","rada wspolnoty",
  "rola rada wspolnoty = agent od razu łączy z opiekunem, bez zadawania pytań"],
], przyklad_pierwszy=False)
for r in (2, 3):
    for c in range(1, 9):
        ws.cell(row=r, column=c).font = PRZYK
        ws.cell(row=r, column=c).fill = ZOLTY

# ============ ZGLOSZENIA ============
ws = wb.create_sheet("zgloszenia")
naglowki(ws, ["kiedy","numer","adres","lokal","kategoria","priorytet","opis","etap","kto_prowadzi",
              "telefon","id_rozmowy","transkrypt","streszczenie","braki","notatka"],
         [18,9,16,10,14,12,44,20,20,16,26,44,34,24,28])
wiersze(ws, [
 ["26.08.2026, 21:47","4472","Sarmacka 12","34","woda","krytyczny",
  "Woda leci z sufitu w łazience, zalewa podłogę, trwa od kilkunastu minut",
  "eskalowane do dyżurnego","dyżurny technik","+48600100200","conv_PRZYKLAD_skasuj",
  "Ada: Dobry wieczór, tu Ada... / Rozmówca: Leje mi się woda z sufitu...",
  "Zgłoszenie zalania w lokalu 34, eskalowane do dyżurnego technika.",
  "","WIERSZ PRZYKŁADOWY, skasuj przed testem"],
], przyklad_pierwszy=True)

# ============ BAZA WIEDZY ============
ws = wb.create_sheet("baza_wiedzy")
naglowki(ws, ["temat","pytanie","odpowiedz_glosowa","dotyczy"], [22,44,86,18])
kb = [
 ["kontakt","Jakie są godziny pracy biura?","Biuro pracuje od poniedziałku do piątku, od dziewiątej do siedemnastej. Poza tymi godzinami przyjmujemy zgłoszenia awarii.","wszystkie"],
 ["kontakt","Jak się z Wami skontaktować?","Telefon dwadzieścia dwa, dwa dziewięć dwa, dziewiętnaście, piętnaście. Adres mailowy to biuro małpa lux myślnik dom kropka pl.","wszystkie"],
 ["platnosci","Do kiedy trzeba zapłacić zaliczkę?","Zaliczka jest płatna do dziesiątego dnia każdego miesiąca, z góry.","wszystkie"],
 ["platnosci","Co składa się na zaliczkę?","Na zaliczkę składają się koszty zarządu nieruchomością wspólną, fundusz remontowy oraz zaliczki na wodę, ścieki i ogrzewanie.","wszystkie"],
 ["platnosci","Jaki jest numer konta?","Numer konta jest w e-kartotece oraz na zawiadomieniu o wysokości zaliczek. Przez telefon podaję tylko cztery ostatnie cyfry do potwierdzenia.","wszystkie"],
 ["platnosci","Ile mam zaległości?","Salda nie podaję przez telefon. Widać je w e-kartotece po zalogowaniu. Mogę też przekazać sprawę do księgowości, która oddzwoni.","wszystkie"],
 ["rozliczenia","Kiedy jest rozliczenie wody i ogrzewania?","Rozliczenie robimy raz w roku, na podstawie odczytów liczników.","wszystkie"],
 ["rozliczenia","Mam nadpłatę, co dalej?","Nadpłata zaliczana jest na poczet przyszłych opłat. Na wniosek właściciela może zostać zwrócona na rachunek.","wszystkie"],
 ["rozliczenia","Mam niedopłatę, czy mogę rozłożyć na raty?","Niedopłatę reguluje się w terminie z rozliczenia. Przy większej kwocie można wystąpić o rozłożenie na raty, wniosek rozpatruje zarząd wspólnoty.","wszystkie"],
 ["rozliczenia","W rozliczeniu jest błąd","Przyjmę to jako zgłoszenie i przekażę do księgowości. Nie oceniam, czy rozliczenie jest prawidłowe.","wszystkie"],
 ["ekartoteka","Co to jest e-kartoteka?","To internetowy portal właściciela lokalu. Są tam salda, historia wpłat, rozliczenia, dokumenty wspólnoty, uchwały i głosowanie elektroniczne.","wszystkie"],
 ["ekartoteka","Nie pamiętam hasła","Na stronie logowania jest opcja odzyskiwania hasła. Link przychodzi na adres mailowy podany w kartotece.","wszystkie"],
 ["ekartoteka","Nie mam dostępu do e-kartoteki","Przyjmę zgłoszenie i przekażę do administratora. Nowy dostęp nadajemy w jeden dzień roboczy.","wszystkie"],
 ["ekartoteka","Czy jest aplikacja na telefon?","Tak, jest aplikacja mobilna. Daje wgląd w te same informacje co portal internetowy.","wszystkie"],
 ["awarie","Co jest awarią krytyczną?","Zapach gazu, dym lub pożar, trwające zalanie, brak wody, prądu albo ogrzewania w całym budynku, uwięzienie w windzie, pęknięty pion, oberwany element elewacji.","wszystkie"],
 ["awarie","Czuję gaz","Proszę natychmiast zadzwonić pod numer sto dwanaście. Ja w tym czasie zgłaszam awarię i łączę z dyżurnym technikiem.","wszystkie"],
 ["odpowiedzialnosc","Cieknie mi kran, kto naprawia?","To instalacja wewnątrz lokalu, czyli po stronie właściciela. Mogę przyjąć zgłoszenie i przekazać administratorowi, ale naprawę zleca i pokrywa właściciel.","wszystkie"],
 ["odpowiedzialnosc","Za co odpowiada wspólnota?","Za części wspólne: piony wodne i kanalizacyjne, instalację elektryczną do licznika, klatki schodowe, dach, elewację, windy, domofony, garaż i tereny zielone.","wszystkie"],
 ["odpowiedzialnosc","Za co odpowiada właściciel?","Za wnętrze swojego lokalu: instalacje za licznikiem i za pierwszym zaworem odcinającym, armaturę, urządzenia, okna od wewnątrz i drzwi wejściowe do lokalu.","wszystkie"],
 ["odpowiedzialnosc","Zalał mnie sąsiad, kto to jest?","Danych innego lokalu nie podaję. Mogę potwierdzić, że sprawa jest zgłoszona, i przekazać ją administratorowi.","wszystkie"],
 ["dokumenty","Potrzebuję zaświadczenia o niezaleganiu","Przygotujemy je do dwóch dni roboczych. Powiem tylko, że jeśli zamawiającym nie jest właściciel lokalu, księgowość poprosi właściciela o zgodę.","wszystkie"],
 ["dokumenty","Ile czeka się na uchwałę albo protokół?","Do trzech dni roboczych. Plan gospodarczy i sprawozdanie finansowe do pięciu dni roboczych.","wszystkie"],
 ["zgody","Czy mogę zabudować balkon?","Tego nie rozstrzygam. Przyjmę wniosek i przekażę do administratora nieruchomości, który odpowie na piśmie.","wszystkie"],
 ["zgody","Chcę zamontować klimatyzację","To wymaga zgody wspólnoty. Przyjmę wniosek i przekażę do administratora, który odpowie na piśmie.","wszystkie"],
 ["zgody","W jakich godzinach można robić remont?","Prace generujące hałas są dopuszczalne w dni robocze, w godzinach od ósmej do dwudziestej.","wszystkie"],
 ["administracja","Zmieniła się liczba osób w lokalu","Przyjmę zgłoszenie telefonicznie. Zmiana wpływa na zaliczki za wodę i odpady, obowiązuje od kolejnego miesiąca.","wszystkie"],
 ["administracja","Kupiłem mieszkanie, co zgłosić?","Prosimy o dane nowego właściciela i stany liczników z dnia przekazania lokalu.","wszystkie"],
 ["administracja","Potrzebuję pilota do bramy","Piloty wydaje administracja po weryfikacji tytułu prawnego do miejsca postojowego. Koszt pokrywa właściciel.","klimczaka 5"],
 ["zebranie","Kiedy jest zebranie roczne?","Zebranie ogółu właścicieli odbywa się raz w roku, do końca marca. Zawiadomienia wysyłamy z co najmniej tygodniowym wyprzedzeniem.","wszystkie"],
 ["zebranie","Czy mogę głosować zdalnie?","Tak, nad uchwałami można głosować elektronicznie, przez e-kartotekę.","wszystkie"],
 ["liczniki","Do kiedy podać odczyt licznika?","Do dziesiątego dnia miesiąca, przez e-kartotekę albo aplikację.","wszystkie"],
]
wiersze(ws, kb)

# ============ SLOWNIKI ============
ws = wb.create_sheet("slowniki")
naglowki(ws, ["pole","narzedzie","dozwolone_wartosci","uwaga"], [24,22,62,44])
wiersze(ws, [
 ["kategoria","utworz_zgloszenie","woda, ogrzewanie, prad, winda, domofon, brama, elewacja, sprzatanie, rozliczenia, inne","przy wątpliwości agent wpisuje inne"],
 ["priorytet","utworz_zgloszenie","krytyczny, zwykly","krytyczny uruchamia natychmiastowe połączenie z dyżurnym"],
 ["akcja","umow_wizyte","pobierz_terminy, potwierdz","najpierw pobierz_terminy, potem potwierdz"],
 ["typ_wizyty","umow_wizyte","technik, odczyt_licznika, odbior_kluczy, spotkanie_z_administratorem",""],
 ["typ_dokumentu","zamow_dokument","zaswiadczenie_o_niezaleganiu, uchwala, protokol_zebrania, plan_gospodarczy, sprawozdanie_finansowe, inne",""],
 ["rola_zamawiajacego","zamow_dokument","wlasciciel, notariusz, posrednik, bank, najemca, inna","rola inna niż wlasciciel wymaga zgody właściciela lokalu"],
 ["forma_odbioru","zamow_dokument","email, osobiscie, poczta","przy email agent prosi o przeliterowanie adresu"],
 ["etap","zgloszenia","przyjęte, eskalowane do dyżurnego, w realizacji, umowione, zamknięte","wypełnia n8n albo człowiek"],
 ["rola","lokale","wlasciciel, najemca, rada wspolnoty, zarzad","rada wspolnoty i zarzad są łączone od razu z człowiekiem"],
])

wb.save("luxdom-baza-agenta.xlsx")
print("zapisane, zakladki:", wb.sheetnames)
